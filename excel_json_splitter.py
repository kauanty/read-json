from __future__ import annotations

import json
import re
import ast
import threading
import traceback
from collections import OrderedDict
from dataclasses import dataclass, field
from pathlib import Path
from queue import Queue, Empty
from typing import Any, Callable, Iterable

from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    import customtkinter as ctk
except ImportError:  # pragma: no cover - exercised manually in restricted hosts
    ctk = None

import tkinter as tk
from tkinter import filedialog, messagebox, ttk


DEFAULT_JSON_COLUMN = "JSON Dados Solicitação"
DEFAULT_SERVICE_TITLE_KEY = "titulo"
DEFAULT_NUMBER_COLUMN = "Número"
EMPTY_VALUE = "Campo Vazio"
JSON_ERROR_VALUE = "ERRO-JSON"
ERROR_SHEET = "ERROS_JSON"
MAX_SHEET_TITLE_LENGTH = 30
MAX_COLUMN_WIDTH = 60
MIN_COLUMN_WIDTH = 12
DATA_ROW_HEIGHT = 15
HEADER_GROUP_COLORS = ("D6EAF8", "C8F2B9")


ProgressCallback = Callable[[int, int], None]


@dataclass
class ProcessConfig:
    input_path: Path
    output_path: Path
    json_column_name: str = DEFAULT_JSON_COLUMN
    service_title_key: str = DEFAULT_SERVICE_TITLE_KEY
    number_column_name: str = DEFAULT_NUMBER_COLUMN


@dataclass
class SheetDataset:
    original_title: str
    root_columns: "OrderedDict[str, OrderedDict[str, None]]" = field(default_factory=OrderedDict)
    rows: list[dict[str, Any]] = field(default_factory=list)

    def add_row(self, number: Any, raw_json: Any, fields: dict[str, OrderedDict[str, Any]]) -> None:
        for root_key, root_fields in fields.items():
            self.root_columns.setdefault(root_key, OrderedDict())
            for column_title in root_fields.keys():
                self.root_columns[root_key].setdefault(column_title, None)

        self.rows.append(
            {
                DEFAULT_NUMBER_COLUMN: sanitize_cell_value(number),
                DEFAULT_JSON_COLUMN: sanitize_cell_value(raw_json),
                "fields": fields,
            }
        )


def sanitize_cell_value(value: Any) -> Any:
    if value is None:
        return EMPTY_VALUE
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return EMPTY_VALUE
        return ILLEGAL_CHARACTERS_RE.sub("", value)
    return value


def is_empty_value(value: Any) -> bool:
    return value is None or (isinstance(value, str) and value.strip() == "")


def normalize_json_text(raw_json: Any) -> str:
    if raw_json is None:
        raise ValueError("JSON vazio")
    text = str(raw_json).strip().lstrip("\ufeff")
    if not text:
        raise ValueError("JSON vazio")
    return text


def parse_json_cell(raw_json: Any) -> Any:
    text = normalize_json_text(raw_json)
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        repaired = repair_json_text(text)
        if repaired != text:
            try:
                return json.loads(repaired)
            except json.JSONDecodeError:
                pass

        try:
            return ast.literal_eval(text)
        except (ValueError, SyntaxError):
            if repaired != text:
                return ast.literal_eval(repaired)
            raise


def repair_json_text(text: str) -> str:
    repaired = text
    repaired = re.sub(r"\bNone\b", "null", repaired)
    repaired = re.sub(r"\bTrue\b", "true", repaired)
    repaired = re.sub(r"\bFalse\b", "false", repaired)
    repaired = re.sub(r",\s*([}\]])", r"\1", repaired)
    repaired = re.sub(
        r'("value"\s*:\s*)(?!["{\[\]}\d\-]|true\b|false\b|null\b)([^,\]}]+)',
        quote_unquoted_json_value,
        repaired,
        flags=re.IGNORECASE,
    )
    return repaired


def quote_unquoted_json_value(match: re.Match[str]) -> str:
    prefix = match.group(1)
    value = match.group(2).strip()
    if not value:
        return f'{prefix}""'
    escaped = value.replace("\\", "\\\\").replace('"', '\\"')
    return f'{prefix}"{escaped}"'


def iter_root_items(parsed_json: Any) -> Iterable[tuple[str, Any]]:
    if isinstance(parsed_json, dict):
        for root_key, root_value in parsed_json.items():
            if isinstance(root_value, list):
                for item in root_value:
                    yield str(root_key), item
            else:
                yield str(root_key), root_value
        return

    if isinstance(parsed_json, list):
        for item in parsed_json:
            yield "JSON", item
        return

    yield "JSON", parsed_json


def extract_item_title(item: Any, service_title_key: str) -> str:
    if isinstance(item, dict):
        title = item.get(service_title_key)
        if not is_empty_value(title):
            return str(title).strip()
        for fallback_key in ("title", "label", "name", "campo"):
            fallback_title = item.get(fallback_key)
            if not is_empty_value(fallback_title):
                return str(fallback_title).strip()
    return "Valor"


def extract_item_value(item: Any) -> Any:
    if isinstance(item, dict):
        for key in ("value", "valor", "Value", "VALOR"):
            if key in item:
                return sanitize_cell_value(item.get(key))
        compact = {key: value for key, value in item.items() if key not in {DEFAULT_SERVICE_TITLE_KEY, "title", "label", "name"}}
        return sanitize_cell_value(json.dumps(compact, ensure_ascii=False)) if compact else EMPTY_VALUE
    return sanitize_cell_value(item)


def flatten_json_fields(parsed_json: Any, service_title_key: str) -> dict[str, OrderedDict[str, Any]]:
    fields: dict[str, OrderedDict[str, Any]] = OrderedDict()
    for root_key, item in iter_root_items(parsed_json):
        root_fields = fields.setdefault(root_key, OrderedDict())
        column_title = extract_item_title(item, service_title_key)
        value = extract_item_value(item)
        root_fields[column_title] = value
    return fields


def choose_service_title(parsed_json: Any, fields: dict[str, OrderedDict[str, Any]], service_title_key: str) -> str:
    if isinstance(parsed_json, dict):
        direct = parsed_json.get(service_title_key)
        if isinstance(direct, str) and direct.strip():
            return direct.strip()

    key_lower = service_title_key.strip().lower()
    for _, item in iter_root_items(parsed_json):
        if isinstance(item, dict):
            for marker_key in ("titulo", "title", "label", "name", "campo"):
                marker = item.get(marker_key)
                if isinstance(marker, str) and marker.strip().lower() == key_lower:
                    item_value = extract_item_value(item)
                    if item_value != EMPTY_VALUE:
                        return str(item_value)

    for root_fields in fields.values():
        for column_title in root_fields.keys():
            if column_title != "Valor":
                return column_title
    return "Sem Título"


def safe_sheet_title(title: str, used_titles: set[str]) -> str:
    cleaned = re.sub(r"[\[\]\:\*\?\/\\]", " ", str(title)).strip() or "Sem Título"
    cleaned = re.sub(r"\s+", " ", cleaned)[:MAX_SHEET_TITLE_LENGTH].strip() or "Sem Título"
    candidate = cleaned
    index = 2
    while candidate in used_titles:
        suffix = f" {index}"
        candidate = f"{cleaned[: MAX_SHEET_TITLE_LENGTH - len(suffix)]}{suffix}".strip()
        index += 1
    used_titles.add(candidate)
    return candidate


def header_fill_for(root_index: int) -> PatternFill:
    color = HEADER_GROUP_COLORS[root_index % len(HEADER_GROUP_COLORS)]
    return PatternFill("solid", fgColor=color)


def read_input_rows(config: ProcessConfig, progress_callback: ProgressCallback | None = None) -> dict[str, SheetDataset]:
    workbook = load_workbook(config.input_path, read_only=True, data_only=True)
    worksheet = workbook.active

    headers = [cell.value for cell in next(worksheet.iter_rows(min_row=1, max_row=1))]
    header_map = {str(header).strip(): index for index, header in enumerate(headers) if header is not None}
    if config.json_column_name not in header_map:
        raise ValueError(f'Coluna JSON "{config.json_column_name}" não encontrada.')
    if config.number_column_name not in header_map:
        raise ValueError(f'Coluna "{config.number_column_name}" não encontrada.')

    json_index = header_map[config.json_column_name]
    number_index = header_map[config.number_column_name]
    total_rows = max(worksheet.max_row - 1, 0)
    datasets: dict[str, SheetDataset] = {}

    for processed, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=1):
        number = row[number_index] if number_index < len(row) else None
        raw_json = row[json_index] if json_index < len(row) else None

        try:
            parsed_json = parse_json_cell(raw_json)
            fields = flatten_json_fields(parsed_json, config.service_title_key)
            service_title = choose_service_title(parsed_json, fields, config.service_title_key)
        except Exception:
            service_title = ERROR_SHEET
            fields = OrderedDict({"Erro": OrderedDict({"Status": JSON_ERROR_VALUE})})

        datasets.setdefault(service_title, SheetDataset(service_title)).add_row(number, raw_json, fields)

        if progress_callback:
            progress_callback(processed, total_rows)

    workbook.close()
    return datasets


def build_output_workbook(datasets: dict[str, SheetDataset], output_path: Path) -> None:
    workbook = Workbook(write_only=False)
    default_sheet = workbook.active
    workbook.remove(default_sheet)

    used_titles: set[str] = set()
    for service_title in sorted(datasets.keys(), key=lambda value: value.casefold()):
        dataset = datasets[service_title]
        worksheet = workbook.create_sheet(safe_sheet_title(service_title, used_titles))
        worksheet.freeze_panes = "A2"

        columns: list[tuple[str, str | None]] = [
            (DEFAULT_NUMBER_COLUMN, None),
            (DEFAULT_JSON_COLUMN, None),
        ]
        for root_key, root_columns in dataset.root_columns.items():
            for column_title in root_columns.keys():
                columns.append((column_title, root_key))

        worksheet.append([column_title for column_title, _ in columns])
        style_header(worksheet, columns)

        for row_data in dataset.rows:
            values = [row_data[DEFAULT_NUMBER_COLUMN], row_data[DEFAULT_JSON_COLUMN]]
            fields = row_data["fields"]
            for column_title, root_key in columns[2:]:
                value = fields.get(root_key, {}).get(column_title, EMPTY_VALUE)
                values.append(sanitize_cell_value(value))
            worksheet.append(values)

        style_data_cells(worksheet)
        auto_adjust_widths(worksheet)

    workbook.save(output_path)
    workbook.close()


def style_header(worksheet: Any, columns: list[tuple[str, str | None]]) -> None:
    root_order: dict[str, int] = {}

    for col_idx, (_, root_key) in enumerate(columns, start=1):
        cell = worksheet.cell(row=1, column=col_idx)
        cell.font = Font(bold=True, color="1F2933")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=False)
        worksheet.row_dimensions[1].height = DATA_ROW_HEIGHT

        if root_key is None:
            cell.fill = PatternFill("solid", fgColor="D6DBDF")
            continue

        root_order.setdefault(root_key, len(root_order))
        cell.fill = header_fill_for(root_order[root_key])


def style_data_cells(worksheet: Any) -> None:
    empty_fill = PatternFill("solid", fgColor="FADBD8")
    empty_font = Font(color="78281F")
    error_font = Font(bold=True, color="FF0000")

    for row in worksheet.iter_rows(min_row=2):
        worksheet.row_dimensions[row[0].row].height = DATA_ROW_HEIGHT
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=False)
            if cell.value == EMPTY_VALUE:
                cell.fill = empty_fill
                cell.font = empty_font
            elif cell.value == JSON_ERROR_VALUE:
                cell.font = error_font


def auto_adjust_widths(worksheet: Any) -> None:
    for header_cell in worksheet[1]:
        column_letter = get_column_letter(header_cell.column)
        header_text = "" if header_cell.value is None else str(header_cell.value)
        width = min(max(len(header_text) + 4, MIN_COLUMN_WIDTH), MAX_COLUMN_WIDTH)
        worksheet.column_dimensions[column_letter].width = width


def process_excel(config: ProcessConfig, progress_callback: ProgressCallback | None = None) -> None:
    datasets = read_input_rows(config, progress_callback)
    build_output_workbook(datasets, config.output_path)


class ExcelJsonSplitterApp:
    def __init__(self) -> None:
        self.ctk_enabled = ctk is not None
        self.queue: Queue[tuple[str, Any]] = Queue()

        if self.ctk_enabled:
            ctk.set_appearance_mode("system")
            ctk.set_default_color_theme("blue")
            self.root = ctk.CTk()
        else:
            self.root = tk.Tk()

        self.root.title("Extrator JSON para Excel")
        self.root.geometry("820x520")
        self.root.minsize(760, 480)

        self.input_path = tk.StringVar(master=self.root)
        self.output_path = tk.StringVar(master=self.root)
        self.json_column_name = tk.StringVar(master=self.root, value=DEFAULT_JSON_COLUMN)
        self.service_title_key = tk.StringVar(master=self.root, value=DEFAULT_SERVICE_TITLE_KEY)
        self.status_text = tk.StringVar(master=self.root, value="Selecione um arquivo Excel para começar.")
        self.progress_text = tk.StringVar(master=self.root, value="0%")

        self._build_ui()
        self.root.after(150, self._poll_queue)

    def _build_ui(self) -> None:
        if self.ctk_enabled:
            frame = ctk.CTkFrame(self.root, corner_radius=8)
            frame.pack(fill="both", expand=True, padx=24, pady=24)
            self._build_ctk_ui(frame)
        else:
            frame = ttk.Frame(self.root, padding=24)
            frame.pack(fill="both", expand=True)
            self._build_tk_ui(frame)

    def _build_ctk_ui(self, frame: Any) -> None:
        ctk.CTkLabel(frame, text="Extrator JSON para Excel", font=ctk.CTkFont(size=24, weight="bold")).pack(anchor="w")
        ctk.CTkLabel(frame, text="Processamento offline com abas por título de serviço.").pack(anchor="w", pady=(4, 20))
        self._ctk_file_row(frame, "Arquivo de entrada", self.input_path, "Selecionar Excel", self.select_input_file)
        self._ctk_file_row(frame, "Arquivo de saída", self.output_path, "Salvar Como...", self.select_output_file)
        self._ctk_entry(frame, "Nome da Coluna JSON", self.json_column_name)
        self._ctk_entry(frame, "Chave do Título do Serviço no JSON", self.service_title_key)
        self.progress_bar = ctk.CTkProgressBar(frame)
        self.progress_bar.set(0)
        self.progress_bar.pack(fill="x", pady=(20, 8))
        ctk.CTkLabel(frame, textvariable=self.progress_text).pack(anchor="w")
        ctk.CTkLabel(frame, textvariable=self.status_text, wraplength=740).pack(anchor="w", pady=(8, 18))
        self.start_button = ctk.CTkButton(frame, text="Iniciar Processamento", command=self.start_processing)
        self.start_button.pack(anchor="e")

    def _ctk_file_row(self, parent: Any, label: str, variable: tk.StringVar, button_text: str, command: Callable[[], None]) -> None:
        ctk.CTkLabel(parent, text=label).pack(anchor="w")
        row = ctk.CTkFrame(parent, fg_color="transparent")
        row.pack(fill="x", pady=(4, 12))
        ctk.CTkEntry(row, textvariable=variable).pack(side="left", fill="x", expand=True, padx=(0, 10))
        ctk.CTkButton(row, text=button_text, width=150, command=command).pack(side="right")

    def _ctk_entry(self, parent: Any, label: str, variable: tk.StringVar) -> None:
        ctk.CTkLabel(parent, text=label).pack(anchor="w")
        ctk.CTkEntry(parent, textvariable=variable).pack(fill="x", pady=(4, 12))

    def _build_tk_ui(self, frame: ttk.Frame) -> None:
        ttk.Label(frame, text="Extrator JSON para Excel", font=("Segoe UI", 18, "bold")).pack(anchor="w")
        ttk.Label(frame, text="Processamento offline com abas por título de serviço.").pack(anchor="w", pady=(4, 20))
        self._tk_file_row(frame, "Arquivo de entrada", self.input_path, "Selecionar Excel", self.select_input_file)
        self._tk_file_row(frame, "Arquivo de saída", self.output_path, "Salvar Como...", self.select_output_file)
        self._tk_entry(frame, "Nome da Coluna JSON", self.json_column_name)
        self._tk_entry(frame, "Chave do Título do Serviço no JSON", self.service_title_key)
        self.progress_bar = ttk.Progressbar(frame, mode="determinate", maximum=100)
        self.progress_bar.pack(fill="x", pady=(20, 8))
        ttk.Label(frame, textvariable=self.progress_text).pack(anchor="w")
        ttk.Label(frame, textvariable=self.status_text, wraplength=740).pack(anchor="w", pady=(8, 18))
        self.start_button = ttk.Button(frame, text="Iniciar Processamento", command=self.start_processing)
        self.start_button.pack(anchor="e")

    def _tk_file_row(self, parent: ttk.Frame, label: str, variable: tk.StringVar, button_text: str, command: Callable[[], None]) -> None:
        ttk.Label(parent, text=label).pack(anchor="w")
        row = ttk.Frame(parent)
        row.pack(fill="x", pady=(4, 12))
        ttk.Entry(row, textvariable=variable).pack(side="left", fill="x", expand=True, padx=(0, 10))
        ttk.Button(row, text=button_text, command=command).pack(side="right")

    def _tk_entry(self, parent: ttk.Frame, label: str, variable: tk.StringVar) -> None:
        ttk.Label(parent, text=label).pack(anchor="w")
        ttk.Entry(parent, textvariable=variable).pack(fill="x", pady=(4, 12))

    def select_input_file(self) -> None:
        selected = filedialog.askopenfilename(filetypes=[("Arquivos Excel", "*.xlsx")])
        if selected:
            self.input_path.set(selected)
            if not self.output_path.get():
                source = Path(selected)
                self.output_path.set(str(source.with_name(f"{source.stem}_processado.xlsx")))

    def select_output_file(self) -> None:
        selected = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Arquivos Excel", "*.xlsx")],
        )
        if selected:
            self.output_path.set(selected)

    def start_processing(self) -> None:
        try:
            config = self._read_config()
        except ValueError as exc:
            messagebox.showwarning("Configuração incompleta", str(exc))
            return

        self._set_processing_state(True)
        self.status_text.set("Processando...")
        worker = threading.Thread(target=self._run_worker, args=(config,), daemon=True)
        worker.start()

    def _read_config(self) -> ProcessConfig:
        input_path = Path(self.input_path.get())
        output_path = Path(self.output_path.get())
        json_column = self.json_column_name.get().strip()
        service_key = self.service_title_key.get().strip()

        if not input_path.is_file():
            raise ValueError("Selecione um arquivo Excel de entrada válido.")
        if input_path.suffix.lower() != ".xlsx":
            raise ValueError("O arquivo de entrada deve estar no formato .xlsx.")
        if not output_path:
            raise ValueError("Escolha onde salvar o arquivo processado.")
        if output_path.suffix.lower() != ".xlsx":
            raise ValueError("O arquivo de saída deve estar no formato .xlsx.")
        if not json_column:
            raise ValueError("Informe o nome da coluna JSON.")
        if not service_key:
            raise ValueError("Informe a chave do título do serviço no JSON.")

        return ProcessConfig(input_path=input_path, output_path=output_path, json_column_name=json_column, service_title_key=service_key)

    def _run_worker(self, config: ProcessConfig) -> None:
        def progress(processed: int, total: int) -> None:
            self.queue.put(("progress", (processed, total)))

        try:
            process_excel(config, progress)
            self.queue.put(("done", str(config.output_path)))
        except Exception as exc:
            self.queue.put(("error", f"{exc}\n\n{traceback.format_exc()}"))

    def _poll_queue(self) -> None:
        try:
            while True:
                event, payload = self.queue.get_nowait()
                if event == "progress":
                    self._update_progress(*payload)
                elif event == "done":
                    self._set_processing_state(False)
                    self.status_text.set(f"Arquivo gerado com sucesso: {payload}")
                    messagebox.showinfo("Processamento concluído", f"Arquivo gerado com sucesso:\n{payload}")
                elif event == "error":
                    self._set_processing_state(False)
                    self.status_text.set("Falha no processamento. Verifique a mensagem de erro.")
                    messagebox.showerror("Erro no processamento", payload)
        except Empty:
            pass
        self.root.after(150, self._poll_queue)

    def _update_progress(self, processed: int, total: int) -> None:
        percent = int((processed / total) * 100) if total else 100
        if self.ctk_enabled:
            self.progress_bar.set(percent / 100)
        else:
            self.progress_bar["value"] = percent
        self.progress_text.set(f"{percent}%")
        self.status_text.set(f"Processando linha {processed:,} de {total:,} ({percent}%)".replace(",", "."))

    def _set_processing_state(self, processing: bool) -> None:
        state = "disabled" if processing else "normal"
        if self.ctk_enabled:
            self.start_button.configure(state=state)
        else:
            self.start_button["state"] = state

    def run(self) -> None:
        self.root.mainloop()


def main() -> None:
    app = ExcelJsonSplitterApp()
    app.run()


if __name__ == "__main__":
    main()
